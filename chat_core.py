"""Explicit-format local import. This module does not access WeChat files."""
import csv
import io
import json
import re
import zipfile
from datetime import datetime, timezone, timedelta
from html import escape

TZ = timezone(timedelta(hours=8))
FIELDS = ('conversation_id', 'conversation', 'sender', 'timestamp', 'type', 'content')


def load_messages(data, filename):
    if len(data) > 20 * 1024 * 1024:
        raise ValueError('文件超过 20 MB，请分批导入。')
    try:
        text = data.decode('utf-8-sig')
        rows = list(csv.DictReader(io.StringIO(text))) if filename.lower().endswith('.csv') else json.loads(text)
    except (UnicodeError, ValueError, csv.Error) as exc:
        raise ValueError('请使用 UTF-8 编码的 CSV 或 JSON 消息数组。') from exc
    if not isinstance(rows, list) or not rows:
        raise ValueError('文件必须包含非空的消息数组或 CSV 数据行。')
    if len(rows) > 100000:
        raise ValueError('单次最多导入 100000 条消息。')
    result = []
    for i, row in enumerate(rows, 1):
        if not isinstance(row, dict) or any(not isinstance(row.get(k), str) for k in FIELDS):
            raise ValueError(f'第 {i} 条记录缺少字段或字段不是字符串：' + ', '.join(FIELDS))
        if any(not row[k].strip() for k in FIELDS[:5]):
            raise ValueError(f'第 {i} 条记录的会话、发送人、类型或时间为空。')
        try:
            dt = datetime.fromisoformat(row['timestamp'].replace('Z', '+00:00'))
            dt = dt.replace(tzinfo=TZ) if dt.tzinfo is None else dt.astimezone(TZ)
        except ValueError as exc:
            raise ValueError(f'第 {i} 条记录的时间无效，请使用 2026-10-20 10:30:00 格式。') from exc
        clean = {k: re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', '', row[k]) for k in FIELDS}
        clean['dt'] = dt
        result.append(clean)
    # Repeated messages may be genuine; do not deduplicate by content.
    return sorted(result, key=lambda m: m['dt'])


def filter_messages(messages, ids, start, end, keyword='', types=None):
    if start > end:
        raise ValueError('开始日期不能晚于结束日期。')
    return [m for m in messages if m['conversation_id'] in ids
            and start <= m['dt'].date() <= end
            and (not keyword or keyword.casefold() in m['content'].casefold())
            and (types is None or m['type'] in types)]


def body(m):
    return m['content'] if m['type'] == 'text' else f"[{m['type']}，附件未嵌入] {m['content']}"


def export_docx(messages, title, note):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(10.5)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    for name in ('Title', 'Heading 1'):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph(title, 'Title')
    doc.add_paragraph(note)
    current = None
    for m in messages:
        if current != m['conversation_id']:
            doc.add_heading(m['conversation'], level=1)
            current = m['conversation_id']
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.add_run(f"{m['dt']:%Y-%m-%d %H:%M:%S}  {m['sender']}").bold = True
        doc.add_paragraph(body(m))
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    section.footer.paragraphs[0]._p.append(field)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def export_pdf(messages, title, note):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    normal = ParagraphStyle('body', fontName='STSong-Light', fontSize=10.5, leading=17, wordWrap='CJK', spaceAfter=8)
    heading = ParagraphStyle('heading', parent=normal, fontSize=15, leading=22, spaceBefore=12, keepWithNext=True)
    meta = ParagraphStyle('meta', parent=normal, fontSize=9, textColor='#555555', keepWithNext=True)
    def paragraph(text, style):
        return Paragraph(escape(text).replace('\n', '<br/>'), style)
    story = [paragraph(title, heading), paragraph(note, normal), Spacer(1, 8)]
    current = None
    for m in messages:
        if current != m['conversation_id']:
            story.append(paragraph(m['conversation'], heading))
            current = m['conversation_id']
        story.append(paragraph(f"{m['dt']:%Y-%m-%d %H:%M:%S}  {m['sender']}", meta))
        story.append(paragraph(body(m) or ' ', normal))
    out = io.BytesIO()
    def page_number(canvas, doc):
        canvas.setFont('STSong-Light', 9)
        canvas.drawCentredString(A4[0] / 2, 25, str(doc.page))
    SimpleDocTemplate(out, pagesize=A4, rightMargin=45, leftMargin=45, topMargin=40, bottomMargin=40).build(story, onFirstPage=page_number, onLaterPages=page_number)
    return out.getvalue()


def export_bundle(messages, start, end, note):
    grouped = {}
    for m in messages:
        grouped.setdefault(m['conversation_id'], []).append(m)
    ordered = [m for group in grouped.values() for m in group]
    docx = export_docx(ordered, '聊天记录导出', note)
    pdf = export_pdf(ordered, '聊天记录导出', note)
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('导出说明.txt', note + '\n非文字消息仅包含说明，图片、语音、视频等原始附件未嵌入。')
        for i, group in enumerate(grouped.values(), 1):
            name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', group[0]['conversation']).strip(' .')[:60] or '会话'
            stem = f'{i:03d}_{name}_{start}_{end}'
            archive.writestr(stem + '.docx', export_docx(group, '聊天记录导出', note))
            archive.writestr(stem + '.pdf', export_pdf(group, '聊天记录导出', note))
    return docx, pdf, out.getvalue()
